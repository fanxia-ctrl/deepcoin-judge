#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""judge.jsonl → 人工复核用的 xlsx。

一行一个 case，后面若干列是**本条命中的判据**（只出命中的，没命中的不占列）。
标注两件事：
  1 每条命中判对了没有 —— 抓过判
  2 有没有该命中却没命中的 —— 抓漏判
先 judge 再标注，所以表里的机器结论是给人核的，不是给人填的。
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from rubric import THEMES, THEME_BY_TIP, TIP_BY_CODE

FONT = "Arial"
MAX_HIT_COLS = 6

JUDGE_FILL = PatternFill("solid", fgColor="EEF3FA")   # 机器判的，只读
MARK_FILL = PatternFill("solid", fgColor="FFF7CC")    # 要人填的
HEAD_FILL = PatternFill("solid", fgColor="2A4B7C")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

OPT_RIGHT = ["对", "错", "拿不准"]
OPT_MISS = ["无", "有", "拿不准"]
OPT_CODES = [f"{t.code} {t.name}" for t in TIP_BY_CODE.values()]


def _hit_cell(v: dict) -> str:
    t = TIP_BY_CODE[v["code"]]
    extra = ""
    if t.type_field and v.get(t.type_field):
        extra = f"（{v[t.type_field]}）"
    return (f"{t.code} {t.name}{extra}\n"
            f"引用：{v.get('quote') or '—'}\n"
            f"理由：{v.get('why') or '—'}")


def build(rows: list[dict], out: Path) -> Path:
    hit_lists = []
    for r in rows:
        vs = [v for v in (r.get("verdicts") or []) if v.get("hit") is True]
        vs.sort(key=lambda v: -TIP_BY_CODE[v["code"]].weight_for(
            str(v.get(TIP_BY_CODE[v["code"]].type_field) or "")))
        hit_lists.append(vs[:MAX_HIT_COLS])
    n_hit = max([len(h) for h in hit_lists] + [1])

    wb = Workbook()
    ws = wb.active
    ws.title = "标注"

    head = ["case_id", "suite", "route_case", "问题", "回答",
            "机器门禁", "扣分", "命中数", "规则命中"]
    for i in range(1, n_hit + 1):
        head += [f"命中{i}", f"命中{i} 判对吗"]
    head += ["有漏判吗", "漏了哪条", "漏判说明", "备注"]
    ws.append(head)

    for r, vs in zip(rows, hit_lists):
        line = [r.get("case_id", ""), r.get("suite", ""), r.get("route_case", ""),
                r.get("query", ""), r.get("answer") or r.get("voice") or "",
                r.get("verdict", ""), r.get("deduct", 0), len(vs),
                "、".join(r.get("contract_hits", []) + r.get("scored_rules", [])) or ""]
        for i in range(n_hit):
            line += [_hit_cell(vs[i]) if i < len(vs) else "", ""]
        line += ["", "", "", ""]
        ws.append(line)

    n_rows = len(rows) + 1
    first_hit = 10
    miss_col = first_hit + n_hit * 2

    for c in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 30

    mark_cols = {first_hit + i * 2 + 1 for i in range(n_hit)} | {
        miss_col, miss_col + 1, miss_col + 2, miss_col + 3}
    for row in ws.iter_rows(min_row=2, max_row=n_rows, max_col=len(head)):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            cell.fill = MARK_FILL if cell.column in mark_cols else JUDGE_FILL
        ws.row_dimensions[row[0].row].height = 96

    widths = {1: 14, 2: 14, 3: 12, 4: 30, 5: 46, 6: 11, 7: 7, 8: 7, 9: 12}
    for i in range(n_hit):
        widths[first_hit + i * 2] = 42
        widths[first_hit + i * 2 + 1] = 11
    widths.update({miss_col: 10, miss_col + 1: 22, miss_col + 2: 26, miss_col + 3: 20})
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "E2"

    # 选项放独立表引用：内联列表有 255 字符上限，长了会被静默丢掉
    opt = wb.create_sheet("_选项")
    for i, v in enumerate(OPT_RIGHT, 1):
        opt.cell(row=i, column=1, value=v)
    for i, v in enumerate(OPT_MISS, 1):
        opt.cell(row=i, column=2, value=v)
    for i, v in enumerate(OPT_CODES, 1):
        opt.cell(row=i, column=3, value=v)
    opt.sheet_state = "hidden"

    def dv(ref: str, col: int, strict: bool = True) -> None:
        d = DataValidation(type="list", formula1=ref, allow_blank=True)
        d.showErrorMessage = strict
        ws.add_data_validation(d)
        L = get_column_letter(col)
        d.add(f"{L}2:{L}{n_rows}")

    for i in range(n_hit):
        dv(f"_选项!$A$1:$A${len(OPT_RIGHT)}", first_hit + i * 2 + 1)
    dv(f"_选项!$B$1:$B${len(OPT_MISS)}", miss_col)
    dv(f"_选项!$C$1:$C${len(OPT_CODES)}", miss_col + 1, strict=False)

    # 判据表：核漏判时对着看
    rf = wb.create_sheet("判据表")
    rf.append(["判据", "主题", "名称", "什么算命中", "前提（条件型才有）", "做对的样子", "权重"])
    for th in THEMES:
        for t in th.tips:
            rf.append([t.code, f"{th.no} {th.name}", t.name, t.desc,
                       t.applies_when or "—", t.right_looks_like or "—", t.weight])
    for c in range(1, 8):
        cell = rf.cell(row=1, column=c)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rf.iter_rows(min_row=2, max_row=rf.max_row, max_col=7):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for c, w in {1: 11, 2: 20, 3: 24, 4: 52, 5: 32, 6: 32, 7: 7}.items():
        rf.column_dimensions[get_column_letter(c)].width = w
    rf.freeze_panes = "A2"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main(judge_jsonl: Path, out: Path | None = None) -> int:
    p = Path(judge_jsonl)
    if p.is_dir():
        p = p / "judge.jsonl"
    if not p.exists():
        print(f"找不到 {p}")
        return 1
    rows = [json.loads(x) for x in p.read_text(encoding="utf-8").splitlines() if x.strip()]
    rows.sort(key=lambda r: (r.get("suite") or "", r.get("case_id") or ""))
    dest = Path(out) if out else p.with_name("复核表.xlsx")
    build(rows, dest)
    n_hit = sum(1 for r in rows for v in (r.get("verdicts") or []) if v.get("hit") is True)
    print(f"OK {dest}\n   {len(rows)} 个 case · {n_hit} 条命中待核 · "
          f"{sum(1 for r in rows if not any(v.get('hit') is True for v in (r.get('verdicts') or [])))} 条零命中（只核漏判）")
    return 0
