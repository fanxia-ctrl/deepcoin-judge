#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""人机一致率 —— judge 好不好用，只看这个。

三层：门禁（能不能当上线卡点）、逐主题 kappa（哪个维度不行）、逐判据过判率。
一致率之外都给 kappa：正负不平衡时准确率会虚高。
"""
from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

from rubric import GRADE_FAIL, HUMAN_ISSUE_TO_TIPS, RATIOS, THEME_BY_TIP, TIP_BY_CODE


def kappa(a: list[bool], b: list[bool]) -> float:
    n = len(a)
    if not n:
        return 0.0
    po = sum(x == y for x, y in zip(a, b)) / n
    pa, pb = sum(a) / n, sum(b) / n
    pe = pa * pb + (1 - pa) * (1 - pb)
    return 0.0 if pe >= 1 else round((po - pe) / (1 - pe), 3)


def prf(tp, fp, fn):
    p = tp / (tp + fp) if tp + fp else 0.0
    r = tp / (tp + fn) if tp + fn else 0.0
    return round(p, 3), round(r, 3), round(2 * p * r / (p + r), 3) if p + r else 0.0


def read_human(xlsx: Path) -> dict[str, dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("需要 openpyxl：pip3 install openpyxl")
    ws = load_workbook(xlsx, data_only=True)["标注"]
    out = {}
    for r in range(4, ws.max_row + 1):
        cid = ws.cell(r, 1).value
        if cid:
            out[str(cid)] = {"grade": str(ws.cell(r, 8).value or "")[:1],
                             "issue": str(ws.cell(r, 9).value or "").strip(),
                             "why": str(ws.cell(r, 10).value or "").strip()}
    return out


def main(judge_jsonl: Path, labels_xlsx: Path, out: Path | None = None) -> int:
    J = {r["case_id"]: r for r in
         (json.loads(l) for l in Path(judge_jsonl).read_text(encoding="utf-8").splitlines()
          if l.strip())}
    H = read_human(Path(labels_xlsx))
    dropped = [c for c in J if not J[c].get("complete", True)]
    common = [c for c in J if c in H and H[c]["grade"] and J[c].get("complete", True)]
    if not common:
        sys.exit("没有可对齐的 case_id")

    L = ["# 人机一致率", "",
         f"judge `{Path(judge_jsonl).parent.name}` · 标注 `{Path(labels_xlsx).name}` · "
         f"可对齐 {len(common)} 条", ""]
    if dropped:
        L += [f"> 排除 **{len(dropped)}** 条判定不完整（缺材料且结论为「干净」，"
              f"或判据组失败）—— 不能拿半份判定算一致率。", ""]

    jf = [J[c]["verdict"] == "不可上线" for c in common]
    hf = [H[c]["grade"] in GRADE_FAIL for c in common]
    tp = sum(x and y for x, y in zip(jf, hf))
    fp = sum(x and not y for x, y in zip(jf, hf))
    fn = sum((not x) and y for x, y in zip(jf, hf))
    tn = sum((not x) and (not y) for x, y in zip(jf, hf))
    p, r, f1 = prf(tp, fp, fn)
    L += ["## 1 门禁一致率", "",
          "judge 判「不可上线」对上人判 C/D。**这是能不能当上线水文测试的门槛。**", "",
          "| | 人判 C/D | 人判 A/B |", "|---|---:|---:|",
          f"| judge 拦下 | {tp} | {fp} |", f"| judge 放过 | {fn} | {tn} |", "",
          f"- 一致率 **{(tp + tn) / len(common):.1%}**　kappa **{kappa(jf, hf)}**",
          f"- 精确率 {p:.1%}　召回率 {r:.1%}　F1 {f1}",
          f"- 漏放 {fn} 条（judge 放过、人判 C/D）—— 水文测试最危险的那类",
          f"- 过拦 {fp} 条 —— 这类多了会让人不信 judge", ""]

    L += ["## 2 逐主题 kappa", "",
          "**优化 judge 按这张表逐个维度做，不要看总分** —— 总分掩盖哪个维度不行。", "",
          "| 主题 | 人判命中 | judge 命中 | 一致率 | kappa |", "|---|---:|---:|---:|---:|"]
    issue_themes = {i: {THEME_BY_TIP[t].key for t in tips if t in THEME_BY_TIP}
                    for i, tips in HUMAN_ISSUE_TO_TIPS.items()}
    for tk in sorted({k for v in issue_themes.values() for k in v}):
        hv = [tk in issue_themes.get(H[c]["issue"], set()) for c in common]
        jv = [any(THEME_BY_TIP[x].key == tk for x in J[c]["neg_hits"]) for c in common]
        L.append(f"| {tk} | {sum(hv)} | {sum(jv)} | "
                 f"{sum(x == y for x, y in zip(hv, jv)) / len(common):.1%} | {kappa(hv, jv)} |")

    L += ["", "## 3 逐判据", "",
          "「命中即 C 的比例」低说明这条在过判 —— 要么描述不清，要么权重给高了。", "",
          "| 判据 | 名称 | 权重 | judge 命中 | 其中人判 C/D | 命中即 C | 标注 n |",
          "|---|---|---:|---:|---:|---:|---:|"]
    by_tip = defaultdict(list)
    for c in common:
        for code in J[c]["neg_hits"]:
            by_tip[code].append(c)
    for code, cases in sorted(by_tip.items(), key=lambda kv: -len(kv[1])):
        t = TIP_BY_CODE[code]
        nc = sum(1 for c in cases if H[c]["grade"] in GRADE_FAIL)
        L.append(f"| `{code}` | {t.name} | {t.weight} | {len(cases)} | {nc} | "
                 f"{nc / len(cases):.0%} | {t.n or '待标注'} |")
    unhit = [c for c in TIP_BY_CODE if c not in by_tip]
    if unhit:
        L += ["", f"一次都没命中：{'、'.join(f'`{c}`' for c in unhit)} —— "
                  "要么这批没有，要么 judge 判不出来，得单独造样本验。", ""]

    L += ["", "## 4 六个比率", "",
          "版本对比用这个，比「加分总和」可解释。", "",
          "| 比率 | 做对 / 该做对 | 达成 |", "|---|---:|---:|"]
    for rn in RATIOS:
        num = sum(J[c]["ratio_num"].get(rn, 0) for c in common)
        den = sum(J[c]["ratio_den"].get(rn, 0) for c in common)
        L.append(f"| {rn} | {num} / {den} | {f'{num / den:.0%}' if den else '—'} |")

    miss = [c for c in common if J[c]["verdict"] != "不可上线" and H[c]["grade"] in GRADE_FAIL]
    if miss:
        L += ["", "## 5 漏放明细", "", "优化 judge 优先看这些。", ""]
        for c in miss[:25]:
            L.append(f"- `{c}`　人判 {H[c]['issue'] or '未填'}　judge {J[c]['verdict']}"
                     f"（命中 {'、'.join(J[c]['neg_hits']) or '无'}）")
            if H[c]["why"]:
                L.append(f"  - 人写的理由：{H[c]['why'][:110]}")
        if len(miss) > 25:
            L.append(f"- …其余 {len(miss) - 25} 条见 judge.jsonl")

    out = Path(out) if out else Path(judge_jsonl).with_name("agreement.md")
    out.write_text("\n".join(L) + "\n", encoding="utf-8")
    print(f"OK {out}")
    print(f"   门禁一致率 {(tp + tn) / len(common):.1%}  kappa {kappa(jf, hf)}  "
          f"漏放 {fn}  过拦 {fp}　（排除不完整 {len(dropped)} 条）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(Path(sys.argv[1]), Path(sys.argv[2]),
                          Path(sys.argv[3]) if len(sys.argv) > 3 else None))
